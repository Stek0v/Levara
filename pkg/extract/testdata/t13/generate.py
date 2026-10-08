"""Author fixtures from frozen source labels; never read extraction output."""
import hashlib
import json
import subprocess
from pathlib import Path
import zipfile
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).parent
FIRST = "First paragraph. Copper count 17."
SECOND = "Второй абзац. Zinc count 23."
CELLS = [["Item", "Units"], ["Copper", "17"], ["Zinc", "23"]]
GOLD = FIRST + " Item Units Copper 17 Zinc 23 " + SECOND
OCR_GOLD = "Copper count 17. Zinc count 23. Total units 40."
# These labels and criteria are authored before any parser/OCR execution.
CASES = []
def add(name, mode="text", text=GOLD, pages=None, paragraphs=None, cells=None):
    case = {"file": name, "mode": mode, "text": text}
    if pages is not None: case["pages"] = pages
    if paragraphs is not None: case["paragraphs"] = paragraphs
    if cells is not None: case["cells"] = cells
    CASES.append(case)

for name in ["ledger.txt", "ledger.pdf", "ledger.docx", "ledger.pptx", "ledger.xlsx", "ledger.html", "ledger.epub", "ledger.odt"]:
    add(name, pages=2 if name.endswith((".pdf", ".pptx")) else None, paragraphs=[FIRST, SECOND], cells=sum(CELLS, []))
add("unicode.txt", text="Привет, команда! 日本語 café 😀\nLine two.", pages=1)
add("empty.txt", "empty", "", 1)
for name in ["corrupt.pdf", "corrupt.docx", "corrupt.epub", "empty.pdf", "encrypted.pdf", "invalid-utf8.txt", "nul.txt"]:
    add(name, "error", "")
add("scanned.pdf", "scan-boundary", "", 1)
add("ocr.png", "ocr", OCR_GOLD, 1)
add("narration.wav", "audio-unavailable", OCR_GOLD)
manifest = {
    "schema_version": 1,
    "authored_before_extraction": True,
    "normalization": "Document equality: ordered Unicode letter/number tokens; punctuation and whitespace are layout separators. Unicode TXT also requires exact original bytes. OCR: whitespace-only normalization, case and punctuation retained.",
    "criteria": {"document_token_sequence": "exact equality (100% retention/order, no extra tokens)", "paragraph_order": "all authored paragraphs in order", "table_cells": "all authored cells in row-major order after the first paragraph", "page_count": "exact where labeled", "ocr_max_cer": 0.05, "ocr_max_wer": 0.10, "audio_quality": "unmeasured: no configured Whisper engine"},
    "cases": CASES,
}
# Write the gold first. Only fixture hashes are filled after generation.
(ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
(ROOT / "ledger.txt").write_text(GOLD, encoding="utf-8")
(ROOT / "unicode.txt").write_text(CASES[8]["text"], encoding="utf-8")
(ROOT / "empty.txt").write_bytes(b"")
(ROOT / "corrupt.pdf").write_bytes(b"%PDF-1.7\nbroken object")
for name in ["corrupt.docx", "corrupt.epub"]: (ROOT / name).write_bytes(b"PK\x03\x04broken archive")
(ROOT / "empty.pdf").write_bytes(b"")
(ROOT / "invalid-utf8.txt").write_bytes(b"\xffinvalid")
(ROOT / "nul.txt").write_bytes(b"valid\x00hidden")

font = "/System/Library/Fonts/Supplemental/Arial.ttf"
pdfmetrics.registerFont(TTFont("T13Arial", font))
c = canvas.Canvas(str(ROOT / "ledger.pdf"), pagesize=(612, 792), invariant=1)
c.setFont("T13Arial", 15); c.drawString(60, 690, FIRST)
for row, values in enumerate(CELLS):
    for col, value in enumerate(values): c.drawString(60 + col*150, 620-row*35, value)
    c.line(55, 610-row*35, 330, 610-row*35)
c.showPage(); c.setFont("T13Arial", 15); c.drawString(60, 690, SECOND); c.save()

d = Document(); d.add_paragraph(FIRST); table = d.add_table(rows=3, cols=2)
for r, values in enumerate(CELLS):
    for col, value in enumerate(values): table.cell(r, col).text = value
d.add_paragraph(SECOND); d.save(ROOT / "ledger.docx")

p = Presentation(); slide=p.slides.add_slide(p.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = FIRST
table=slide.shapes.add_table(3,2,Inches(1),Inches(2),Inches(5),Inches(2)).table
for r, values in enumerate(CELLS):
    for col, value in enumerate(values): table.cell(r,col).text=value
slide=p.slides.add_slide(p.slide_layouts[6]); slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text=SECOND
p.save(ROOT / "ledger.pptx")

w=Workbook(); sheet=w.active; sheet.title="Ledger"; sheet.append([FIRST])
for values in CELLS: sheet.append(values)
sheet.append([SECOND]); w.save(ROOT / "ledger.xlsx")

table_html = "<table>" + "".join("<tr>"+"".join("<td>"+v+"</td>" for v in row)+"</tr>" for row in CELLS)+"</table>"
html="<html><body><p>"+FIRST+"</p>"+table_html+"<p>"+SECOND+"</p></body></html>"
(ROOT / "ledger.html").write_text(html,encoding="utf-8")
with zipfile.ZipFile(ROOT / "ledger.epub","w") as z:
    z.writestr("mimetype","application/epub+zip",compress_type=zipfile.ZIP_STORED)
    z.writestr("META-INF/container.xml",'<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    z.writestr("OEBPS/content.opf",'<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="book">t13-ledger</dc:identifier><dc:title>Ledger</dc:title><dc:language>en</dc:language></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>')
    z.writestr("OEBPS/chapter.xhtml",html.replace("<html>",'<html xmlns="http://www.w3.org/1999/xhtml">'))
with zipfile.ZipFile(ROOT / "ledger.odt","w") as z:
    z.writestr("mimetype","application/vnd.oasis.opendocument.text",compress_type=zipfile.ZIP_STORED)
    z.writestr("META-INF/manifest.xml",'<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    rows="".join('<table:table-row>'+"".join('<table:table-cell office:value-type="string"><text:p>'+escape(v)+'</text:p></table:table-cell>' for v in row)+'</table:table-row>' for row in CELLS)
    z.writestr("content.xml",'<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" office:version="1.2"><office:body><office:text><text:p>'+FIRST+'</text:p><table:table table:name="Ledger">'+rows+'</table:table><text:p>'+SECOND+'</text:p></office:text></office:body></office:document-content>')

im=Image.new("RGB",(1000,420),"white"); draw=ImageDraw.Draw(im); image_font=ImageFont.truetype(font,54)
for y,line in enumerate(["Copper count 17.","Zinc count 23.","Total units 40."]): draw.text((75,60+y*100),line,font=image_font,fill="black")
im.save(ROOT / "ocr.png")
c=canvas.Canvas(str(ROOT / "scanned.pdf"),pagesize=(612,792),invariant=1); c.drawImage(str(ROOT / "ocr.png"),50,440,width=510,height=214);c.save()
encrypted=PdfWriter(); encrypted.append(PdfReader(ROOT / "ledger.pdf")); encrypted.encrypt("t13-reader"); encrypted.write(ROOT / "encrypted.pdf")

subprocess.run(["say", "-v", "Fred", "-o", str(ROOT / "narration.aiff"), OCR_GOLD],check=True,timeout=30)
subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-i", str(ROOT / "narration.aiff"), "-ar", "16000", "-ac", "1", "-c:a", "pcm_s16le", str(ROOT / "narration.wav")],check=True,timeout=30)
(ROOT / "narration.aiff").unlink()

for case in CASES: case["sha256"]=hashlib.sha256((ROOT / case["file"]).read_bytes()).hexdigest()
(ROOT / "manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n")
print("Authored labels and criteria frozen; fixtures generated without parser/OCR output.")
